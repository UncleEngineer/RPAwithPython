from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

import os

pythonfile_path = os.getcwd()
excel_folder = os.path.join(pythonfile_path,'excelfile')
file_list = os.listdir(excel_folder)

for f in file_list:
    excel_name = f
    excel_path = os.path.join(excel_folder, excel_name)

    excelfile = load_workbook(excel_path)
    sheet = excelfile.active

    # ADD GRAPH TO CURRENT FILE
    values = Reference(sheet,min_col=3,max_col=3, min_row=2,max_row=4)
    chart = BarChart()
    chart.title = 'ความสูงต้นไม้'
    chart.y_axis.title = 'ความสูง'
    chart.x_axis.title = 'ชนิดต้นไม้'
    chart.legend = None
    chart.add_data(values)
    sheet.add_chart(chart,'E5')

    cat = Reference(sheet, min_col=1,max_col=1, min_row=2, max_row=4)
    chart.set_categories(cat)

    excelfile.save(excel_path)